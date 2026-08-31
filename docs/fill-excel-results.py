"""Fill Status / Actual Result / Executed At / Method into FieldPro_UI_Test_Cases.xlsx for every row."""
import json
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(r"e:\ashhadu\FieldWorker_System\docs")
XLSX = ROOT / "FieldPro_UI_Test_Cases.xlsx"
RESULTS = ROOT / "ALL_TC_RESULTS.json"

STATUS_FILLS = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "BLOCKED": PatternFill("solid", fgColor="FFEB9C"),
    "SKIP": PatternFill("solid", fgColor="D9D9D9"),
}
STATUS_FONT = {
    "PASS": Font(color="006100", bold=True),
    "FAIL": Font(color="9C0006", bold=True),
    "BLOCKED": Font(color="9C5700", bold=True),
}

def main():
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    results = data["results"]
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Test Cases"]

    # Ensure header columns L-O
    headers = {
        12: "Status",
        13: "Actual Result / Notes",
        14: "Executed At",
        15: "Method",
    }
    for col, title in headers.items():
        cell = ws.cell(1, col, title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(bold=True, color="FFFFFF")

    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    missing = []
    for row in range(2, ws.max_row + 1):
        tc = ws.cell(row, 1).value
        if not tc:
            continue
        tc = str(tc).strip()
        res = results.get(tc)
        if not res:
            missing.append(tc)
            res = {
                "status": "BLOCKED",
                "notes": "No result in runner output",
                "at": data.get("ranAt"),
                "method": "pending",
            }
        status = res.get("status", "BLOCKED")
        notes = res.get("notes", "")
        executed = res.get("at") or data.get("ranAt") or datetime.utcnow().isoformat()
        method = res.get("method", "auto")

        c_status = ws.cell(row, 12, status)
        c_notes = ws.cell(row, 13, notes)
        c_at = ws.cell(row, 14, executed)
        c_method = ws.cell(row, 15, method)

        for c in (c_status, c_notes, c_at, c_method):
            c.border = thin
            c.alignment = Alignment(wrap_text=True, vertical="top")

        if status in STATUS_FILLS:
            c_status.fill = STATUS_FILLS[status]
            c_status.font = STATUS_FONT.get(status, Font(bold=True))

    ws.column_dimensions["L"].width = 12
    ws.column_dimensions["M"].width = 55
    ws.column_dimensions["N"].width = 24
    ws.column_dimensions["O"].width = 12

    # Update Summary sheet with live counts
    summary = data.get("summary", {})
    if "Summary" in wb.sheetnames:
        sm = wb["Summary"]
        # append a results block at the bottom
        start = sm.max_row + 2
        sm.cell(start, 1, "LIVE EXECUTION RESULTS").font = Font(bold=True, size=14)
        sm.cell(start + 1, 1, f"FE: {data.get('fe')}")
        sm.cell(start + 2, 1, f"API: {data.get('api')}")
        sm.cell(start + 3, 1, f"Ran at: {data.get('ranAt')}")
        sm.cell(start + 5, 1, "Status")
        sm.cell(start + 5, 2, "Count")
        sm.cell(start + 5, 1).font = Font(bold=True)
        sm.cell(start + 5, 2).font = Font(bold=True)
        r = start + 6
        for k in ("PASS", "FAIL", "BLOCKED", "SKIP"):
            sm.cell(r, 1, k)
            sm.cell(r, 2, summary.get(k, 0))
            if k in STATUS_FILLS:
                sm.cell(r, 1).fill = STATUS_FILLS[k]
            r += 1
        sm.cell(r + 1, 1, "Note: BLOCKED = mutating/email/Twilio/Stripe/multi-step cases not auto-run on shared live DB.")
        sm.cell(r + 2, 1, f"Missing from runner map: {len(missing)}")

    out = ROOT / "FieldPro_UI_Test_Cases_RESULTS.xlsx"
    wb.save(out)
    # also overwrite original so user sees filled cells in the given file
    wb.save(XLSX)
    print(f"Saved {out}")
    print(f"Updated {XLSX}")
    print("Summary:", summary)
    print("Missing TCs:", missing[:20], "..." if len(missing) > 20 else "")

if __name__ == "__main__":
    main()
